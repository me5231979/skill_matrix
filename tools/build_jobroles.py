#!/usr/bin/env python3
"""Join the Oracle org-structure export to the SBJA skill framework.

Staff know their JOB_ROLE ("Student Advisor IC2"), not their sub-family code
("AcadAdv"). This builds the lookup that turns a job role into the SBJA
sub-family + career level whose skill profile the app already carries.

Output: assets/data/jobroles.json
"""
import json, re, sys, collections
import openpyxl

ORG = sys.argv[1]
SBJA = sys.argv[2]
OUT = sys.argv[3]

# Career levels: the export uses SS2/SS3/SS4 where the framework uses S2/S3/S4.
LEVEL_FIX = {'SS1': 'S1', 'SS2': 'S2', 'SS3': 'S3', 'SS4': 'S4'}
# Levels that carry no framework proficiency (legacy/unleveled job codes).
UNLEVELED = {'IC', 'PM', 'FM', 'None', ''}

# Sub-family codes the abbreviation matcher gets wrong or cannot see.
# Left: "FAMILY|CODE" from the export. Right: SBJA sub-family, or None for
# "no profile exists" (Executive, Temporary Services and friends).
OVERRIDE = {
    'Development & Alumni Relations|DonorRelEvent': 'Stewardship/Donor Relations',
    'Student Services|IntSvcs': 'International Services',   # sits in Academic Affairs in SBJA
    'Public Safety|SafeTech': 'Campus Safety & Security',
    'Information Technology|BI': 'Data Infrastructure',
    'Information Technology|None': None,
    'Information Technology|ITMan': 'IT Consulting, Collaboration, & Relationship Management',
    'Research|None': None,
    'People Operations|HR': 'Consulting',
    'People Operations|None': None,
    'Financial Operations|None': None,
    'Health Services|None': None,
    'Internal & External Relations|None': None,
    'Campus Support Services|None': None,
    'Educational & Academic Support|None': None,
    'Engineering & Architecture|None': None,
    'Executive|None': None,
    'Legal & Compliance|None': None,
    'Public & Environmental Safety|None': None,
    'Unassigned|None': None,
}


def norm(s):
    return re.sub(r'[^a-z]', '', str(s).lower())


def abbrev_score(code, name):
    """How well does an abbreviation like 'CuInstDe' fit 'Curriculum & Instructional Development'?"""
    a, n = norm(code), norm(name)
    if a == n:
        return 100
    initials = ''.join(w[0] for w in re.findall(r'[A-Za-z]+', name))
    if a == norm(initials):
        return 95
    # every letter of the code must appear in order; contiguous runs score higher
    i, contiguous = 0, 0
    for ch in a:
        j = n.find(ch, i)
        if j < 0:
            return 0
        if j == i:
            contiguous += 1
        i = j + 1
    return 50 + contiguous


def load_sbja(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    by_family = collections.defaultdict(set)
    all_subs = set()
    for r in wb['Job Family Skills'].iter_rows(min_row=6, values_only=True):
        if r[1] and r[2]:
            by_family[str(r[1]).strip()].add(str(r[2]).strip())
            all_subs.add(str(r[2]).strip())
    wb.close()
    return by_family, all_subs


def main():
    by_family, all_subs = load_sbja(SBJA)

    wb = openpyxl.load_workbook(ORG, read_only=True)
    rows = [r for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)
            if r and str(r[5]) == 'A' and r[4]]
    wb.close()

    resolved = {}       # "FAMILY|CODE" -> sbja sub-family or None
    for r in rows:
        fam, code = str(r[0]), str(r[1])
        key = fam + '|' + code
        if key in resolved:
            continue
        if key in OVERRIDE:
            resolved[key] = OVERRIDE[key]
            continue
        best, best_score = None, 0
        for cand in by_family.get(fam, ()):
            s = abbrev_score(code, cand)
            if s > best_score:
                best, best_score = cand, s
        resolved[key] = best if best_score >= 50 else None

    roles, seen = [], set()
    for r in rows:
        fam, code, lvl, jobcode, name = str(r[0]), str(r[1]), str(r[2]), r[3], str(r[4]).strip()
        lvl = LEVEL_FIX.get(lvl, lvl)
        if lvl in UNLEVELED:
            lvl = ''
        key = (name, lvl)
        if key in seen:
            continue
        seen.add(key)
        roles.append({
            'n': name,
            'c': jobcode,
            'f': fam,
            's': resolved.get(fam + '|' + code),
            'l': lvl,
            'h': int(r[9] or 0),
        })

    roles.sort(key=lambda x: x['n'])
    profiled = [x for x in roles if x['s']]
    hc_all = sum(x['h'] for x in roles)
    hc_prof = sum(x['h'] for x in profiled)

    json.dump({'roles': roles}, open(OUT, 'w'), separators=(',', ':'))

    print(f'{len(roles)} active job roles -> {OUT}')
    print(f'  with a skill profile: {len(profiled)} roles, {hc_prof} of {hc_all} headcount '
          f'({100 * hc_prof / hc_all:.1f}%)')
    unprofiled = collections.defaultdict(int)
    for x in roles:
        if not x['s']:
            unprofiled[x['f']] += x['h']
    print('  no profile, by family (headcount):')
    for f, h in sorted(unprofiled.items(), key=lambda kv: -kv[1]):
        print(f'    {f}: {h}')
    missing_lvl = [x for x in roles if not x['l']]
    print(f'  roles with no career level: {len(missing_lvl)} '
          f'({sum(x["h"] for x in missing_lvl)} headcount)')


if __name__ == '__main__':
    main()
