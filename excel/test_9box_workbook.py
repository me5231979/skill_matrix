#!/usr/bin/env python3
"""Fill the worked example into a copy of the template and verify every automated result."""
import datetime, shutil, subprocess, sys, json
import openpyxl

import os
SP = os.path.dirname(os.path.abspath(__file__))
SRC = f"{SP}/VU_9Box_Talent_Assessment_Template.xlsx"
TST = f"{SP}/test_filled.xlsx"
# Recalculation needs LibreOffice; point RECALC_SCRIPT at a helper that opens,
# recalculates, and re-saves the file (or open/save the two test files in Excel
# manually and re-run with SKIP_RECALC=1).
RECALC = os.environ.get("RECALC_SCRIPT", "/root/.claude/skills/synced/xlsx/scripts/recalc.py")

shutil.copy(SRC, TST)
wb = openpyxl.load_workbook(TST)
st, rs, a, sk = wb["Setup"], wb["1. Role Scoring"], wb["2. Assessment"], wb["3. Skills (SBJA)"]

st["B4"] = "People Operations"
st["B5"] = "EC Example"
st["B6"] = "Chair Example"
st["B8"] = datetime.date(2026, 11, 6)
st["B9"] = "Pilot"

roles = [
    ("People Operations Executive", "E1", 5, 4, 4, 4, 4, None, None, "Core function impaired in weeks; rare expertise; thin bench."),
    ("Total Rewards Lead", "M5", 4, 3, 4, 3, 4, None, None, "Hot comp market, scarce strategy expertise, thin bench."),
    ("Consulting Lead", "M5", 3, 4, 3, 2, 3, "Yes", "Held at Not critical", "Exceptional practice variability; credible bench and stable market."),
    ("Learning & Development Lead", "M4", 3, 3, 3, 2, 3, None, None, "Important and visible, but refillable."),
    ("Payroll Lead", "M4", 4, 2, 2, 2, 2, None, None, "Touches every employee; processes documented; market refills the seat."),
]
for i, (t, lv, f, g, h, ib, mk, disc, outc, rat) in enumerate(roles):
    r = 6 + i
    rs[f"B{r}"] = t; rs[f"C{r}"] = lv
    rs[f"F{r}"], rs[f"G{r}"], rs[f"H{r}"], rs[f"I{r}"], rs[f"J{r}"] = f, g, h, ib, mk
    if disc: rs[f"M{r}"] = disc
    if outc: rs[f"N{r}"] = outc
    rs[f"P{r}"] = rat

people = [
    ("J. Rivera", "People Operations Executive", 48, "No", 2, datetime.date(2026, 3, 15), 3,
     "Engagement up 9 points; $1.1M saved; led SBJA rollout.", "High", "High", "High",
     "Accept proposed", None, None, "1-2 yrs", "No", None, "Green", "Executive coaching"),
    ("T. Chen", "Learning & Development Lead", 60, "No", 2, datetime.date(2026, 4, 1), 3,
     "Two academies launched on time/under budget; completion +18%.", "High", "Medium", "High",
     "Accept proposed", None, None, None, None, None, None, "Leadership Development gap plan"),
    ("R. Dubois", "Consulting Lead", 72, "No", 0, datetime.date(2026, 5, 1), 3,
     "Client satisfaction highest in function; redesigned intake; trains consultants.", "Medium", "Low", "High",
     "Accept proposed", None, None, None, None, None, None, "Recognize and formalize teaching role"),
    ("S. Patel", "Total Rewards Lead", 30, "No", 1, datetime.date(2026, 6, 1), 2,
     "Clean open enrollment; renewal under budget; one delayed vendor audit.", "High", "High", "High",
     "Accept proposed", None, None, "1-2 yrs", "Yes", "Critical role + external market heat; retention action within 2 weeks.", "Yellow", "Executive-track development plan"),
    ("L. Novak", "Payroll Lead", 40, "No", 0, datetime.date(2026, 2, 1), 1,
     "Missed accuracy targets two quarters; recency probe surfaced full year.", "Medium", "Medium", "Medium",
     "Support pathway - no label", None, "Panel routed to support pathway; coaching plan outside this program.", None, None, None, None, None),
]
cols = ["B", "C", "H", "I", "J", "K", "M", "N", "P", "Q", "R", "W", "X", "Z", "AA", "AB", "AC", "AD", "AF"]
for i, vals in enumerate(people):
    r = 6 + i
    for c, v in zip(cols, vals):
        if v is not None:
            a[f"{c}{r}"] = v

skills = [
    ("T. Chen", "Instructional Design", "Expert", "Expert"),
    ("T. Chen", "Curriculum Development", "Expert", "Expert"),
    ("T. Chen", "Adult Learning Principles", "Expert", "Expert"),
    ("T. Chen", "Training Analysis", "Advanced", "Advanced"),
    ("T. Chen", "Learning Platforms", "Intermediate", "Advanced"),
    ("T. Chen", "Leadership Development", "Advanced", "Intermediate"),
    ("J. Rivera", "Executive Leadership", "Expert", "Expert"),
    ("J. Rivera", "Org Design", "Advanced", "Expert"),
    ("S. Patel", "Compensation Strategy", "Expert", "Advanced"),
    ("S. Patel", "Benefits Design", "Advanced", "Advanced"),
]
for i, (p, s_, req, cur) in enumerate(skills):
    r = 6 + i
    sk[f"A{r}"], sk[f"C{r}"], sk[f"D{r}"], sk[f"E{r}"] = p, s_, req, cur

wb.save(TST)
out = subprocess.run([sys.executable, RECALC, TST, "300"], capture_output=True, text=True)
res = json.loads(out.stdout)
assert res.get("status") == "success" and res.get("total_errors") == 0, res

wb = openpyxl.load_workbook(TST, data_only=True)
rs, a, g, d, ex = wb["1. Role Scoring"], wb["2. Assessment"], wb["9-Box Grid"], wb["Dashboard"], wb["Export (flat)"]

fails = []
def chk(label, got, want):
    ok = (got == want)
    if not ok:
        fails.append(f"FAIL {label}: got {got!r}, want {want!r}")
    print(("PASS" if ok else "FAIL"), label, "->", repr(got))

def chk_contains(label, got, sub):
    ok = isinstance(got, str) and sub in got
    if not ok:
        fails.append(f"FAIL {label}: got {got!r}, want contains {sub!r}")
    print(("PASS" if ok else "FAIL"), label, "->", repr(got))

# Step 1: totals and critical calls (document appendix)
for r, tot in zip(range(6, 11), [21, 18, 15, 14, 12]):
    chk(f"Role total row {r}", rs[f"K{r}"].value, tot)
for r, o in zip(range(6, 11), ["YES", "YES", "No", "No", "No"]):
    chk(f"CRITICAL? row {r}", rs[f"O{r}"].value, o)
chk_contains("Consulting borderline flag", rs["L8"].value, "borderline")
chk("Total Rewards (18) NOT borderline", "borderline" in str(rs["L7"].value), False)
chk("L&D (14) NOT borderline", "borderline" in str(rs["L9"].value), False)
chk("Role QA all clean", [rs[f"Q{r}"].value for r in range(6, 11)], [None] * 5)

# Step 2: potential, boxes (document appendix)
for r, p in zip(range(6, 11), ["High", "Medium", "Low", "High", "Medium"]):
    chk(f"Potential row {r}", a[f"S{r}"].value, p)
want_boxes = ["High Potential", "Emerging Talent", "High Professional",
              "Solid Performer, Strong Potential", "Support pathway (no label)"]
for r, b in zip(range(6, 11), want_boxes):
    chk(f"Final box row {r}", a[f"Y{r}"].value, b)
chk("Novak proposed box", a["T10"].value, "Inconsistent Performer")
for r, c in zip(range(6, 11), ["YES", "No", "No", "YES", "No"]):
    chk(f"Critical lookup row {r}", a[f"F{r}"].value, c)
chk("Chen skills at/above 83%", round(a["U7"].value, 4), round(5 / 6, 4))
chk("Chen gap count", a["V7"].value, 1)
chk("Dubois retrospective flag", a["O8"].value, "Retrospective - extra probe")
chk("Rivera flags complete", a["AE6"].value, "Complete")
chk("Chen flags n/a", a["AE7"].value, "n/a - not critical")
chk("Assessment QA all clean", [a[f"AH{r}"].value for r in range(6, 11)], [None] * 5)

# Grid
chk_contains("Grid (3,High) names", g["D6"].value, "J. Rivera")
chk_contains("Grid (3,Med) names", g["C6"].value, "T. Chen")
chk_contains("Grid (3,Low) names", g["B6"].value, "R. Dubois")
chk_contains("Grid (2,High) names", g["D8"].value, "S. Patel")
chk_contains("Grid HiPo count", g["D5"].value, "(1)")
chk("Grid people entered", g["C13"].value, 5)
chk("Grid placed", g["C14"].value, 4)
chk("Grid support pathway", g["C17"].value, 1)
chk("Grid top-box share 50%", g["C19"].value, 0.5)
chk_contains("Grid concentration warning", g["C20"].value, "Concentrated result")

# Dashboard
chk("Output 1 status", d["D6"].value, "Complete")
chk("Output 2 status", d["D7"].value, "Complete")
chk("Output 3 status", d["D8"].value, "Complete")
chk("Output 4 status", d["D9"].value, "Ready")
chk("Unit status", d["D10"].value, "UNIT COMPLETE ✓")
chk("Retrospective count", d["C25"].value, 2)
chk("Coverage Green", d["C39"].value, 1)
chk("Coverage Yellow", d["C40"].value, 1)
chk("Retention flags", d["C45"].value, 1)
chk_contains("Retention list", d["B48"].value, "S. Patel")

# Export
chk("Export unit", ex["A4"].value, "People Operations")
chk("Export person", ex["D4"].value, "J. Rivera")
chk("Export final box", ex["Q4"].value, "High Potential")
chk("Export empty row blank", ex["D9"].value, None)

# ---- negative pass: break things, expect QA to catch them ----
wb2 = openpyxl.load_workbook(TST)
wb2["1. Role Scoring"]["P10"] = None            # rationale missing
wb2["2. Assessment"]["AC9"] = None              # retention action missing (Patel)
wb2["2. Assessment"]["K7"] = datetime.date(2024, 1, 1)  # stale career conversation (Chen)
TST2 = f"{SP}/test_broken.xlsx"
wb2.save(TST2)
out = subprocess.run([sys.executable, RECALC, TST2, "300"], capture_output=True, text=True)
res = json.loads(out.stdout)
assert res.get("status") == "success" and res.get("total_errors") == 0, res
wb2 = openpyxl.load_workbook(TST2, data_only=True)
rs2, a2, d2 = wb2["1. Role Scoring"], wb2["2. Assessment"], wb2["Dashboard"]
chk_contains("QA catches missing rationale", rs2["Q10"].value, "Rationale missing")
chk_contains("QA catches missing risk action", a2["AH9"].value, "Retention action missing")
chk_contains("QA catches stale conversation", a2["AH7"].value, "Career conversation needed")
chk("Patel flags show risk action missing", a2["AE9"].value, "Risk action missing")
chk("Output 4 flips to Issues open", d2["D9"].value, "Issues open")
chk("Unit status flips", d2["D10"].value, "IN PROGRESS")

print()
if fails:
    print(f"{len(fails)} FAILURES")
    for f in fails: print(" ", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
